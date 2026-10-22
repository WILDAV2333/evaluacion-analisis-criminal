import streamlit as st
import pandas as pd
import sqlite3
import json
import io
import os
from datetime import datetime
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side

st.set_page_config(
    page_title="Evaluación - Proyecto Integrador",
    page_icon="⚖️",
    layout="wide",
    initial_sidebar_state="collapsed"
)

# -------------------------------------------------------------
# CONFIGURACIÓN MAESTRA PRECARGADA (NUNCA SE BORRA ANTE INACTIVIDAD)
# -------------------------------------------------------------
DATOS_PRECONFIGURADOS = {
    "Paralelo A": {
        "docentes": [
            "Sgos. Freddy Andrés Loza Barrionuevo",
            "Dr. Galo Alexander Quiñones Escuntar",
            "Dra. Nora Alexandra Mantilla Allan",
            "Mayor Patricio Ramos (Suplente)"
        ],
        "grupos": {
            "Grupo Nro. 1": {
                "tema": "CASO: KEVIN ZAMBRANO",
                "alumnos": [
                    "Cptn. Morales Torres Elvis Santiago",
                    "Sgos. Guaño Paguay Diego Armando",
                    "Cbop. Mina Pastrana Yadira Tatiana",
                    "Cbos. Escobar Munoz Fernando Patricio",
                    "Cbos. Pindo Jiménez Antony Cristhofer"
                    "Poli. Paredes Arcos Marco Ruben"
                ]
            },
            "Grupo Nro. 2": {
                "tema": "CASO: LETY CANDO",
                "alumnos": [
                    "Cptn. Torres Chasi Jaime Fernando",
                    "Cbop. Zanafria Montoya Gonzalo Daniel",
                    "Cbop. Yanza Chasiquiza Diego Fernando",
                    "Cbop. Flores Cárdenas Luis Felipe"
                    "Cbos. Palacios Jara Cristian Joel"
                ]
            },
            "Grupo Nro. 3": {
                "tema": "CASO: MAFIA ALBANEZA",
                "alumnos": [
                    "Tnte. Calan Lozada Erick Alexander",
                    "Cbop. Coyago Narvaez Santiago Paúl",
                    "Cbop. Calero Lara Mauricio Israel",
                    "Cbos. Tenorio Ávila Andrea Mishell"
                    "Cbos. Minda Cuamacas Gabriel Alexander"
                    "Poli. León Jacome Juan Carlos"
                ]
            },
            "Grupo Nro. 4": {
                "tema": "CASO: ANTIOQUEÑO II",
                "alumnos": [
                    "Cptn. Garzón Masapanta Henry Rodrigo",
                    "Sgos. Cupueran Vallejos Juan Diego",
                    "Cbos. Torres Palacios Darwin Paul",
                    "Poli. Melo Macias Evelyn Micaela"
                    "Poli. Pasquel Jaramillo Luis Alfredo"
                ]
            },
            "Grupo Nro. 5": {
                "tema": "CASO: BRIONES CHIQUITO",
                "alumnos": [
                    "Tnte. Silva Toledo Michael Eduardo",
                    "Sgop. Casa Caiza German Hernan",
                    "Cbop. Ortiz Medina Douglas Enrique",
                    "Cbos. Sánchez Torres Lorena Magaly"
                    "Poli. López Puente Dylan Issac"
                    "Poli. Mejia Rigcha Jorge Alberto"
                ]
            }
        }
    },
    "Paralelo B": {
        "docentes": [
            "Mayr. Mauro Napoleon Gavilanez Ortega",
            "Dra. Silvana Alexandra Sandoval Meléndez",
            "Dra. Elisa Margarita Fiallos Montaleza",
            "Mayor Patricio Ramos (Suplente)"
        ],
        "grupos": {
            "Grupo Nro. 1": {
                "tema": "CASO: PAOLA ORTEGA,
                "alumnos": [
                    "Tnte. Bocancho Velasque Freddy Marcelo",
                    "Cbos. Picuasi Diaz Roberto Fabian",
                    "Cbos. Quilumba Ponce Kevin Marcelo",
                    "Poli. Aguilar Cordova Luis Jandry",
                    "Poli. Chango Gorotiza Maria Jose"
                ]
            },
            "Grupo Nro. 2": {
                "tema": "CASO: ANTIOQUIA",
                "alumnos": [
                    "Sbte. Perez Bonilla Edyan Steven",
                    "Sgos. Zapata Navas Christian Vinicio",
                    "Cbop. Arias Shiguango Angel Stalin",
                    "Cbop. Quintana Pabon Marco Andres",
                    "Poli. Flores Osorio Alex Roberto"
                ]
            },
            "Grupo Nro. 3": {
                "tema": "CASO: NICOLS",
                "alumnos": [
                    "Cptn. Montes Oquendo Edgar Rafael",
                    "Sgos. Cajas Paredes Edwin Mauricio",
                    "Cbop. Cobeña Echeverria Jonathan Andres",
                    "Cbos. Bastidas Feijoo Judith Michelle",
                    "Cbos. Lanchi Chamba Luis Eduardo"
                ]
            },
            "Grupo Nro. 4": {
                "tema": "CASO: CHRISTIAN GUAÑUNA",
                "alumnos": [
                    "Sbte. Logacho Sanchez Eduardo Andres",
                    "Cbos. Granda Granda Karla Ibeth",
                    "Cbos. Jaramillo Vera Steven Hoover",
                    "Poli. Amaguaya Damian Danny David",
                    "Poli. Pilco Yugsan Jose Ivan"
                ]
            },
            "Grupo Nro. 5": {
                "tema": "CASO: DIEGO CABEZAS",
                "alumnos": [
                    "Cptn. Guaman Cabascango Henrry Eduardo",
                    "Cbop. Quinancela Martillo Edison Geovanny",
                    "Cbos. Jimenez Parrales Briggitte Ginger",
                    "Cbos. Narvaez Chalan Marlon Stalin",
                    "Poli. Chancusig Mendoza Henry Ismael"
                    "Poli. Quishpe Jitala Alexis Adrian"
                ]
            }
        }
    }
}

ARCHIVO_RESPALDO = "respaldo_evaluaciones.json"

def get_db():
    conn = sqlite3.connect("calificaciones_defensa.db", check_same_thread=False)
    cursor = conn.cursor()
    
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS evaluaciones (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            fecha TEXT,
            paralelo TEXT,
            evaluador TEXT,
            grupo TEXT,
            alumno TEXT,
            nota_escrito REAL,
            detalles_escrito_json TEXT,
            oral_sustentacion_nivel TEXT,
            oral_sustentacion_pts REAL,
            oral_interrogatorio_nivel TEXT,
            oral_interrogatorio_pts REAL,
            oral_estrategico_nivel TEXT,
            oral_estrategico_pts REAL,
            oral_diferenciador_nivel TEXT,
            oral_diferenciador_pts REAL,
            nota_oral REAL,
            nota_total REAL
        )
    """)
    
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS config_docentes (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            paralelo TEXT,
            nombre TEXT
        )
    """)
    
    cursor.execute("""
        CREATE TABLE IF NOT EXISTS config_grupos (
            id INTEGER PRIMARY KEY AUTOINCREMENT,
            paralelo TEXT,
            grupo TEXT,
            tema_caso TEXT,
            alumnos_json TEXT
        )
    """)
    
    cursor.execute("SELECT COUNT(*) FROM config_docentes")
    if cursor.fetchone()[0] == 0:
        for par, data in DATOS_PRECONFIGURADOS.items():
            for doc in data["docentes"]:
                cursor.execute("INSERT INTO config_docentes (paralelo, nombre) VALUES (?, ?)", (par, doc))
            for grp_nom, g_info in data["grupos"].items():
                cursor.execute(
                    "INSERT INTO config_grupos (paralelo, grupo, tema_caso, alumnos_json) VALUES (?, ?, ?, ?)",
                    (par, grp_nom, g_info["tema"], json.dumps(g_info["alumnos"]))
                )
    
    cursor.execute("SELECT COUNT(*) FROM evaluaciones")
    if cursor.fetchone()[0] == 0 and os.path.exists(ARCHIVO_RESPALDO):
        try:
            with open(ARCHIVO_RESPALDO, "r", encoding="utf-8") as f:
                respaldo = json.load(f)
                for reg in respaldo:
                    cursor.execute("""
                        INSERT INTO evaluaciones 
                        (fecha, paralelo, evaluador, grupo, alumno, nota_escrito, detalles_escrito_json,
                         oral_sustentacion_nivel, oral_sustentacion_pts,
                         oral_interrogatorio_nivel, oral_interrogatorio_pts,
                         oral_estrategico_nivel, oral_estrategico_pts,
                         oral_diferenciador_nivel, oral_diferenciador_pts,
                         nota_oral, nota_total)
                        VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                    """, reg)
        except Exception:
            pass

    conn.commit()
    return conn

conn = get_db()

def respaldar_evaluaciones_en_disco():
    cursor = conn.cursor()
    cursor.execute("""
        SELECT fecha, paralelo, evaluador, grupo, alumno, nota_escrito, detalles_escrito_json,
               oral_sustentacion_nivel, oral_sustentacion_pts,
               oral_interrogatorio_nivel, oral_interrogatorio_pts,
               oral_estrategico_nivel, oral_estrategico_pts,
               oral_diferenciador_nivel, oral_diferenciador_pts,
               nota_oral, nota_total FROM evaluaciones
    """)
    filas = cursor.fetchall()
    try:
        with open(ARCHIVO_RESPALDO, "w", encoding="utf-8") as f:
            json.dump(filas, f, ensure_ascii=False)
    except Exception:
        pass

def cargar_configuracion(paralelo):
    cursor = conn.cursor()
    cursor.execute("SELECT nombre FROM config_docentes WHERE paralelo = ? ORDER BY id ASC", (paralelo,))
    docentes = [row[0] for row in cursor.fetchall()]
    
    cursor.execute("SELECT grupo, alumnos_json, tema_caso FROM config_grupos WHERE paralelo = ? ORDER BY id ASC", (paralelo,))
    grupos_db = cursor.fetchall()
    
    grupos_dict = {}
    temas_dict = {}
    for g_nombre, alums_raw, tema in grupos_db:
        try:
            grupos_dict[g_nombre] = json.loads(alums_raw)
        except Exception:
            grupos_dict[g_nombre] = []
        temas_dict[g_nombre] = tema if tema else "Caso Práctico Integrador"
            
    return docentes, grupos_dict, temas_dict

def separar_grado_nombre(cadena_alumno):
    partes = cadena_alumno.strip().split(maxsplit=1)
    if not partes:
        return "S/G", ""
    
    primer_token = partes[0].strip()
    token_lower = primer_token.lower()
    
    mapa_grados = {
        "cptn.": "Cptn.", "cptn": "Cptn.",
        "tnte.": "Tnte.", "tnte": "Tnte.",
        "sbte.": "Sbte.", "sbte": "Sbte.", "subt.": "Sbte.", "subt": "Sbte.",
        "sgop.": "Sgop.", "sgop": "Sgop.", "sgto.": "Sgop.", "sgto": "Sgop.",
        "sgos.": "Sgos.", "sgos": "Sgos.",
        "cbop.": "Cbop.", "cbop": "Cbop.",
        "cbos.": "Cbos.", "cbos": "Cbos.",
        "poli.": "Poli.", "poli": "Poli.", "policia": "Poli.", "policía": "Poli.",
        "myor.": "Myor.", "myor": "Myor.", "mayor": "Myor.",
        "crnl.": "Crnl.", "crnl": "Crnl.", "coronel": "Crnl.",
        "gral.": "Gral.", "gral": "Gral."
    }
    
    if token_lower in mapa_grados:
        grado_oficial = mapa_grados[token_lower]
        resto_nombres = partes[1].strip() if len(partes) > 1 else ""
        return grado_oficial, resto_nombres
    elif (token_lower.rstrip(".") + ".") in mapa_grados:
        grado_oficial = mapa_grados[token_lower.rstrip(".") + "."]
        resto_nombres = partes[1].strip() if len(partes) > 1 else ""
        return grado_oficial, resto_nombres
        
    return "S/G", cadena_alumno.strip()

OPCIONES_NIVEL = ["Excelente", "Bueno", "Básico", "Deficiente"]
PARALELOS_DISPONIBLES = ["Paralelo A", "Paralelo B"]

CRITERIOS_DETALLADOS_ESCRITO = [
    {
        "criterio": "ENFOQUE LEGAL Y TIPIFICACIÓN",
        "total_pts": 1.00,
        "indicadores": [
            {
                "desc": "Identificación clara del tipo penal aplicable al caso; sustento en normativa legal vigente (COIP u otra normativa pertinente).",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Relación coherente entre los hechos analizados y su tipificación jurídica.",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Argumentación legal estructurada (elementos del delito: conducta, tipicidad, antijuridicidad, culpabilidad).",
                "max": 0.34,
                "valores": {"Excelente": 0.34, "Bueno": 0.25, "Básico": 0.16, "Deficiente": 0.01}
            }
        ]
    },
    {
        "criterio": "COMPRENSIÓN CRIMINOLÓGICA E HIPÓTESIS",
        "total_pts": 1.00,
        "indicadores": [
            {
                "desc": "Habilidades para la formulación de hipótesis investigativas claras y verificables.",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Identificación de patrones delictivos (modus operandi, temporalidad, recurrencia).",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Análisis del fenómeno desde una perspectiva criminológica (causas, motivaciones, comportamiento) y la relación entre variables (víctima, victimario, entorno).",
                "max": 0.34,
                "valores": {"Excelente": 0.34, "Bueno": 0.25, "Básico": 0.16, "Deficiente": 0.01}
            }
        ]
    },
    {
        "criterio": "ANÁLISIS DE INVESTIGACIÓN CRIMINAL",
        "total_pts": 1.00,
        "indicadores": [
            {
                "desc": "Análisis del plan de investigación o presentan un plan de investigación.",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Realizan el mapeo de jerarquías (tipo de estructura criminal, nodos y enlaces (facilitadores)).",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Evaluación procedimental (robustez de las pruebas o evidencias presentadas).",
                "max": 0.34,
                "valores": {"Excelente": 0.34, "Bueno": 0.25, "Básico": 0.16, "Deficiente": 0.01}
            }
        ]
    },
    {
        "criterio": "ANÁLISIS GEOESPACIAL Y TÁCTICO",
        "total_pts": 1.50,
        "indicadores": [
            {
                "desc": "Capacidad de recopilar, procesar, depurar y analizar información proveniente de diversas fuentes.",
                "max": 0.50,
                "valores": {"Excelente": 0.50, "Bueno": 0.38, "Básico": 0.25, "Deficiente": 0.03}
            },
            {
                "desc": "Uso de herramientas tecnológicas y metodologías de análisis de datos.",
                "max": 0.50,
                "valores": {"Excelente": 0.50, "Bueno": 0.38, "Básico": 0.25, "Deficiente": 0.03}
            },
            {
                "desc": "Orientadas a la generación de insumos analíticos que aporten valor al proceso investigativo.",
                "max": 0.50,
                "valores": {"Excelente": 0.50, "Bueno": 0.38, "Básico": 0.25, "Deficiente": 0.02}
            }
        ]
    },
    {
        "criterio": "OSINT Y CIBERINTELIGENCIA",
        "total_pts": 1.00,
        "indicadores": [
            {
                "desc": "Uso de fuentes abiertas (redes sociales, web, registros públicos).",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Aplicación de técnicas básicas de búsqueda, filtrado, verificación y validación de la información obtenida.",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Integración de hallazgos OSINT dentro del análisis general plasmado en el informe.",
                "max": 0.34,
                "valores": {"Excelente": 0.34, "Bueno": 0.25, "Básico": 0.16, "Deficiente": 0.01}
            }
        ]
    },
    {
        "criterio": "PRODUCCIÓN DE INTELIGENCIA Y PROPUESTA",
        "total_pts": 1.50,
        "indicadores": [
            {
                "desc": "Elaboración de un producto de inteligencia estructurado, que contenga una síntesis clara de hallazgos relevantes.",
                "max": 0.50,
                "valores": {"Excelente": 0.50, "Bueno": 0.38, "Básico": 0.25, "Deficiente": 0.03}
            },
            {
                "desc": "Generación de recomendaciones operativas y/o estratégicas.",
                "max": 0.50,
                "valores": {"Excelente": 0.50, "Bueno": 0.38, "Básico": 0.25, "Deficiente": 0.03}
            },
            {
                "desc": "Mantenimiento de un enfoque práctico orientado a la toma de decisiones institucionales.",
                "max": 0.50,
                "valores": {"Excelente": 0.50, "Bueno": 0.38, "Básico": 0.25, "Deficiente": 0.02}
            }
        ]
    },
    {
        "criterio": "PRESENTACIÓN FORMAL DEL DOCUMENTO",
        "total_pts": 1.00,
        "indicadores": [
            {
                "desc": "Estructura lógica del documento entregable físico (inicio, teoría de caso, desarrollo, conclusiones).",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Redacción clara, coherente y sin errores ortográficos.",
                "max": 0.33,
                "valores": {"Excelente": 0.33, "Bueno": 0.25, "Básico": 0.17, "Deficiente": 0.02}
            },
            {
                "desc": "Uso adecuado de tablas, gráficos, mapas e imágenes integradas al informe.",
                "max": 0.34,
                "valores": {"Excelente": 0.34, "Bueno": 0.25, "Básico": 0.16, "Deficiente": 0.01}
            }
        ]
    }
]

CRITERIOS_ORAL = [
    ("Sustentación (Exposición y Defensa Oral)", 3.00, "Dominio técnico del tema expuesto, claridad, seguridad, argumentación sólida y material audiovisual de apoyo."),
    ("Interrogatorio y Resolución ante el Tribunal", 3.00, "Respuestas precisas, fundamentadas, justificación metodológica, aplomo y seguridad profesional."),
    ("Análisis Estratégico y Prospectivo (Defensa de Escenarios)", 3.00, "Identificación de tendencias delictivas, proyección de escenarios futuros y formulación de alertas."),
    ("Elemento Diferenciador (Valor Agregado y Pensamiento Crítico)", 3.00, "Innovación, perspicacia analítica, perspectiva crítica, evitar sesgos y criterio propio.")
]

PUNTOS_ORAL = {
    "Excelente": 3.00,
    "Bueno": 2.25,
    "Básico": 1.50,
    "Deficiente": 0.15
}

def generar_excel_cuadro_individual(df_cuadro, nombre_grupo, paralelo):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Consolidado"
    
    headers = [
        "Ord.", "Grado", "Nombres y Apellidos", 
        "Sustentación (3.00)", "Interrogatorio (3.00)", 
        "Análisis Estrat. (3.00)", "Elem. Diferenc. (3.00)", 
        "TOTAL ORAL (12.00)", "TRABAJO ESCRITO (8.00)", "CALIFICACIÓN FINAL (20.00)"
    ]
    ws.append(headers)
    
    fill_header = PatternFill(start_color="E2EFDA", end_color="E2EFDA", fill_type="solid")
    font_header = Font(name="Arial", size=10, bold=True)
    font_bold = Font(name="Arial", size=10, bold=True)
    font_regular = Font(name="Arial", size=10, bold=False)
    align_center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    align_left = Alignment(horizontal="left", vertical="center")
    
    thin_border = Border(
        left=Side(style='thin', color='000000'),
        right=Side(style='thin', color='000000'),
        top=Side(style='thin', color='000000'),
        bottom=Side(style='thin', color='000000')
    )
    
    for col_num in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_num)
        cell.fill = fill_header
        cell.font = font_header
        cell.alignment = align_center
        cell.border = thin_border
        
    ws.row_dimensions[1].height = 28
    
    for idx, row in df_cuadro.iterrows():
        ws.append([
            row["Ord."],
            row["Grado"],
            row["Nombres y Apellidos"],
            round(float(row["Sustentación (3.00)"]), 2),
            round(float(row["Interrogatorio (3.00)"]), 2),
            round(float(row["Análisis Estrat. (3.00)"]), 2),
            round(float(row["Elem. Diferenc. (3.00)"]), 2),
            round(float(row["TOTAL ORAL (12.00)"]), 2),
            round(float(row["TRABAJO ESCRITO (8.00)"]), 2),
            round(float(row["CALIFICACIÓN FINAL (20.00)"]), 2)
        ])
        curr_r = ws.max_row
        ws.row_dimensions[curr_r].height = 20
        for col_num in range(1, 11):
            c = ws.cell(row=curr_r, column=col_num)
            c.border = thin_border
            c.font = font_bold if col_num in [8, 9, 10] else font_regular
            c.alignment = align_left if col_num == 3 else align_center
            
    widths = [6, 10, 36, 16, 16, 16, 16, 16, 18, 22]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[openpyxl.utils.get_column_letter(i)].width = w
        
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()

def generar_acta_docx_oficial(plantilla_path, paralelo, grupo_nombre, tema_caso, fecha_str, integrantes, tribunal_nombres, notas_escrito_21, desglose_alumnos):
    doc = docx.Document(plantilla_path)
    
    # 1. Encabezado
    doc.paragraphs[5].text = f"Integrantes del Grupo Nro: {grupo_nombre}                           Paralelo: {paralelo}"
    
    for i in range(6):
        p_idx = 6 + i
        if i < len(integrantes):
            doc.paragraphs[p_idx].text = f"{i+1}. {integrantes[i]}"
        else:
            doc.paragraphs[p_idx].text = f"{i+1}. ----------------------------------------"
            
    doc.paragraphs[12].text = f"Tema del Caso Analizado: {tema_caso}"
    doc.paragraphs[13].text = f"Fecha: {fecha_str}"
    
    for i in range(3):
        p_idx = 15 + i
        if i < len(tribunal_nombres):
            doc.paragraphs[p_idx].text = f"{i+1}. {tribunal_nombres[i]}"
        else:
            doc.paragraphs[p_idx].text = f"{i+1}. ----------------------------------------"
            
    # 2. Tabla 0: Fase I (Entregable Físico - 21 Indicadores)
    t0 = doc.tables[0]
    t0.rows[11].cells[1].text = "ANÁLISIS DE INVESTIGACIÓN CRIMINAL"
    t0.rows[12].cells[1].text = "Análisis del plan de investigación o presentan un plan de investigación"
    t0.rows[13].cells[1].text = "Realizan el mapeo de jerarquías (tipo de estructura criminal, nodos y enlaces (facilitadores))"
    t0.rows[14].cells[1].text = "Evaluación procedimental (robustez de las pruebas o evidencias presentadas)"
    t0.rows[21].cells[1].text = "OSINT Y CIBERINTELIGENCIA"
    
    mapa_crit_t0 = [
        ([2, 3, 4], 5),
        ([7, 8, 9], 10),
        ([12, 13, 14], 15),
        ([17, 18, 19], 20),
        ([22, 23, 24], 25),
        ([27, 28, 29], 30),
        ([32, 33, 34], 35)
    ]
    
    tot_escrito_acum = 0.0
    ind_counter = 0
    for ind_rows, tot_row in mapa_crit_t0:
        subtot_crit = 0.0
        for r in ind_rows:
            val = float(notas_escrito_21[ind_counter]) if ind_counter < len(notas_escrito_21) else 0.0
            t0.rows[r].cells[2].text = f"{val:.2f}"
            subtot_crit += val
            ind_counter += 1
        t0.rows[tot_row].cells[2].text = f"{subtot_crit:.2f}"
        tot_escrito_acum += subtot_crit
        
    t0.rows[36].cells[2].text = f"{tot_escrito_acum:.2f} / 8 puntos"
    
    # 3. Eliminar la antigua Tabla 1 y Tabla 2
    t1_old = doc.tables[1]
    t2_old = doc.tables[2]
    t3_firmas = doc.tables[3]
    
    t1_old._element.getparent().remove(t1_old._element)
    t2_old._element.getparent().remove(t2_old._element)
    
    # 4. Insertar encabezado oficial de la FASE II
    p_head = doc.add_paragraph()
    p_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_head.paragraph_format.space_before = Pt(14)
    p_head.paragraph_format.space_after = Pt(6)
    r_head = p_head.add_run("FASE II: SUSTENTACIÓN ORAL, DEFENSA E INTERROGATORIO (Sobre 12 Puntos)")
    r_head.bold = True
    r_head.font.size = Pt(10)
    r_head.font.name = "Arial"
    t3_firmas._element.addprevious(p_head._element)
    
    # 5. Insertar la nueva Tabla Individual con las 10 columnas
    columnas_indiv = [
        "Ord.", "Grado", "Nombres y\nApellidos", 
        "Sustentación\n(3.00)", "Interrogatorio\n(3.00)", 
        "Análisis Estrat.\n(3.00)", "Elem. Diferenc.\n(3.00)", 
        "TOTAL ORAL\n(12.00)", "TRABAJO ESCRITO\n(8.00)", "CALIFICACIÓN FINAL\n(20.00)"
    ]
    
    num_filas = len(desglose_alumnos) + 1
    t_indiv = doc.add_table(rows=num_filas, cols=10)
    t_indiv.style = "Table Grid"
    t_indiv.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    col_widths = [0.4, 0.6, 1.8, 0.6, 0.65, 0.6, 0.6, 0.65, 0.65, 0.75]
    
    hdr_cells = t_indiv.rows[0].cells
    for c_i, col_name in enumerate(columnas_indiv):
        hdr_cells[c_i].text = col_name.replace("\\n", "\n")
        shading = parse_xml(r'<w:shd {} w:fill="E2EFDA"/>'.format(nsdecls('w')))
        hdr_cells[c_i]._tc.get_or_add_tcPr().append(shading)
        p = hdr_cells[c_i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in p.runs:
            run.font.bold = True
            run.font.size = Pt(7.5)
            run.font.name = "Arial"
            
    for row_i, al in enumerate(desglose_alumnos, start=1):
        row_cells = t_indiv.rows[row_i].cells
        grado, nombres = separar_grado_nombre(al["nombre"])
        
        datos_fila = [
            str(row_i),
            grado,
            nombres,
            f"{al['sust']:.2f}",
            f"{al['interrog']:.2f}",
            f"{al['estrateg']:.2f}",
            f"{al['diferenc']:.2f}",
            f"{al['total_oral']:.2f}",
            f"{tot_escrito_acum:.2f}",
            f"{(al['total_oral'] + tot_escrito_acum):.2f}"
        ]
        
        for c_i, valor in enumerate(datos_fila):
            row_cells[c_i].text = valor
            p = row_cells[c_i].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT if c_i == 2 else WD_ALIGN_PARAGRAPH.CENTER
            for run in p.runs:
                run.font.size = Pt(8)
                run.font.name = "Arial"
                if c_i in [7, 8, 9]:
                    run.font.bold = True
                    
    for row in t_indiv.rows:
        for idx, width in enumerate(col_widths):
            row.cells[idx].width = Inches(width)
            
    t3_firmas._element.addprevious(t_indiv._element)
    
    # 6. Tabla de Firmas (Garantizar los 3 evaluadores completos)
    firma_base = "\n\n\n\n___________________________\nFirma: Evaluador "
    t3_firmas.rows[0].cells[0].text = f"{firma_base}1\nNombres: {tribunal_nombres[0] if len(tribunal_nombres) > 0 else '………………………'}"
    t3_firmas.rows[0].cells[1].text = f"{firma_base}2\nNombres: {tribunal_nombres[1] if len(tribunal_nombres) > 1 else '………………………'}"
    t3_firmas.rows[1].cells[0].text = f"{firma_base}3\nNombres: {tribunal_nombres[2] if len(tribunal_nombres) > 2 else '………………………'}"
    
    for row in t3_firmas.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(8.5)
                    run.font.name = "Arial"

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()

def calcular_cuadro_consolidado_grupo(paralelo, grupo_nombre, grupos_p):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT evaluador, alumno, detalles_escrito_json, 
               oral_sustentacion_pts, oral_interrogatorio_pts, 
               oral_estrategico_pts, oral_diferenciador_pts, 
               nota_oral
        FROM evaluaciones
        WHERE paralelo = ? AND grupo = ?
    """, (paralelo, grupo_nombre))
    regs = cursor.fetchall()
    if not regs:
        return None, 0, []
    
    evaluadores_registrados = list(dict.fromkeys(r[0] for r in regs))
    cant_evaluadores = len(evaluadores_registrados)
    
    listas_ind = []
    for r in regs:
        if r[2]:
            try:
                listas_ind.append(json.loads(r[2]))
            except Exception:
                pass
    tot_escrito = 0.0
    if listas_ind:
        num_l = len(listas_ind)
        tot_escrito = round(sum(sum(l) for l in listas_ind) / num_l, 2)
        
    integrantes = grupos_p.get(grupo_nombre, [])
    if not integrantes and regs:
        integrantes = list(dict.fromkeys(r[1] for r in regs))
        
    filas = []
    for idx, al_nom in enumerate(integrantes, start=1):
        evals_al = [r for r in regs if r[1] == al_nom]
        if evals_al:
            sust = round(sum(r[3] for r in evals_al) / len(evals_al), 2)
            interrog = round(sum(r[4] for r in evals_al) / len(evals_al), 2)
            estrateg = round(sum(r[5] for r in evals_al) / len(evals_al), 2)
            diferenc = round(sum(r[6] for r in evals_al) / len(evals_al), 2)
            total_oral = round(sum(r[7] for r in evals_al) / len(evals_al), 2)
        else:
            sust = interrog = estrateg = diferenc = total_oral = 0.0
            
        grado, nombres = separar_grado_nombre(al_nom)
        nota_final = round(total_oral + tot_escrito, 2)
        filas.append({
            "Ord.": idx,
            "Grado": grado,
            "Nombres y Apellidos": nombres,
            "Sustentación (3.00)": sust,
            "Interrogatorio (3.00)": interrog,
            "Análisis Estrat. (3.00)": estrateg,
            "Elem. Diferenc. (3.00)": diferenc,
            "TOTAL ORAL (12.00)": total_oral,
            "TRABAJO ESCRITO (8.00)": tot_escrito,
            "CALIFICACIÓN FINAL (20.00)": nota_final
        })
    df_cuadro = pd.DataFrame(filas)
    return df_cuadro, cant_evaluadores, evaluadores_registrados

def obtener_supletorios_pendientes(paralelo, grupos_p):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT grupo, alumno, AVG(nota_total) as promedio_final
        FROM evaluaciones
        WHERE paralelo = ? AND grupo NOT LIKE 'Supletorio%'
        GROUP BY grupo, alumno
        HAVING promedio_final < 14.00
    """, (paralelo,))
    filas = cursor.fetchall()
    
    supletorios = {}
    for grp, al, prom in filas:
        nom_suplet = f"Supletorio {grp}"
        if nom_suplet not in supletorios:
            supletorios[nom_suplet] = {
                "tema": f"Examen Supletorio de Sustentación Oral - {grp}",
                "alumnos": []
            }
        supletorios[nom_suplet]["alumnos"].append(al)
    return supletorios

def obtener_nota_escrito_grupo_original(paralelo, grupo_nombre):
    cursor = conn.cursor()
    cursor.execute("""
        SELECT nota_escrito, detalles_escrito_json FROM evaluaciones
        WHERE paralelo = ? AND grupo = ? LIMIT 1
    """, (paralelo, grupo_nombre))
    fila = cursor.fetchone()
    if fila:
        return fila[0], fila[1]
    return 0.0, "[]"

# -------------------------------------------------------------
# 3. INTERFAZ: PESTAÑAS
# -------------------------------------------------------------
pestana_evaluar, pestana_admin = st.tabs(["📝 Formulario de Calificación", "🔒 Administrador"])

# =============================================================
# --- PESTAÑA EVALUADOR ---
# =============================================================
with pestana_evaluar:
    banner_candidatos = [
        "banner.jpeg",
        "banner.jpg",
        "banner.png",
        "Banner - II Curso de Esp Ana Criminal.jpeg"
    ]
    banner_encontrado = None
    for b in banner_candidatos:
        if os.path.exists(b):
            banner_encontrado = b
            break
            
    if banner_encontrado:
        st.image(banner_encontrado, use_container_width=True)
    else:
        st.title("Defensa de Proyecto Integrador")
        st.subheader("II Curso de Especialización en Análisis Criminal")
        
    st.markdown("<h3 style='text-align: center; margin-top: 10px; margin-bottom: 25px;'>Defensa de Proyecto Integrador</h3>", unsafe_allow_html=True)
    
    paralelo_sel = st.selectbox(
        "Seleccione el Paralelo:",
        PARALELOS_DISPONIBLES,
        index=0
    )
    
    col_eval, col_grp = st.columns(2)
    evaluador_activo = None
    grupo_activo = None
    alumnos_activos = []
    tema_activo = ""
    docentes_paralelo = []
    grupos_paralelo = {}
    temas_paralelo = {}
    es_supletorio = False
    
    if paralelo_sel:
        docentes_paralelo, grupos_paralelo, temas_paralelo = cargar_configuracion(paralelo_sel)
        
        with col_eval:
            if docentes_paralelo:
                evaluador_activo = st.selectbox(
                    f"Miembro del Tribunal ({paralelo_sel}):",
                    docentes_paralelo,
                    index=None,
                    placeholder="Seleccione su nombre..."
                )
            else:
                st.warning(f"No hay evaluadores asignados al {paralelo_sel}.")
                
        with col_grp:
            if grupos_paralelo and evaluador_activo:
                cursor = conn.cursor()
                cursor.execute("""
                    SELECT DISTINCT grupo FROM evaluaciones 
                    WHERE paralelo = ? AND evaluador = ?
                """, (paralelo_sel, evaluador_activo))
                grupos_ya_calificados = [r[0] for r in cursor.fetchall()]
                
                grupos_ordinarios_pendientes = [g for g in grupos_paralelo.keys() if g not in grupos_ya_calificados]
                
                supletorios_dict = obtener_supletorios_pendientes(paralelo_sel, grupos_paralelo)
                grupos_supletorio_pendientes = [s for s in supletorios_dict.keys() if s not in grupos_ya_calificados]
                
                opciones_grupo = grupos_ordinarios_pendientes + grupos_supletorio_pendientes
                
                if opciones_grupo:
                    grupo_activo = st.selectbox(
                        "Grupo a Evaluar (Pendientes de su calificación):",
                        opciones_grupo,
                        index=None,
                        placeholder="Seleccione el grupo a calificar..."
                    )
                    if grupo_activo:
                        if grupo_activo in grupos_paralelo:
                            es_supletorio = False
                            alumnos_activos = grupos_paralelo[grupo_activo]
                            tema_activo = temas_paralelo.get(grupo_activo, "Caso Práctico Integrador")
                        elif grupo_activo in supletorios_dict:
                            es_supletorio = True
                            alumnos_activos = supletorios_dict[grupo_activo]["alumnos"]
                            tema_activo = supletorios_dict[grupo_activo]["tema"]
                else:
                    st.success(f"👏 ¡Estimado/a {evaluador_activo}, usted ha completado todas las evaluaciones asignadas en el {paralelo_sel}!")
                    
                if grupos_ya_calificados:
                    st.caption(f"🔒 Evaluaciones ya registradas por usted: {', '.join(grupos_ya_calificados)}")
            elif not evaluador_activo:
                st.info("👆 Seleccione primero su nombre como Miembro del Tribunal.")

    # Módulo de transparencia y consulta consolidada
    if paralelo_sel and grupos_paralelo:
        with st.expander("🔍 Ver y Descargar Cuadro Consolidado Oficial (Transparencia)"):
            st.caption("Consulte las calificaciones de los grupos evaluados por el tribunal.")
            supletorios_act = obtener_supletorios_pendientes(paralelo_sel, grupos_paralelo)
            todos_grupos_consulta = list(grupos_paralelo.keys()) + list(supletorios_act.keys())
            grp_consultar = st.selectbox("Seleccione el Grupo para consultar:", todos_grupos_consulta, key="consultar_transp_grp")
            if grp_consultar:
                df_cuadro_grp, cant_evals, lista_evals = calcular_cuadro_consolidado_grupo(paralelo_sel, grp_consultar, grupos_paralelo)
                total_tribunal_necesario = 3
                
                if df_cuadro_grp is not None and not df_cuadro_grp.empty:
                    if cant_evals >= total_tribunal_necesario:
                        st.success(f"✅ ¡GRUPO COMPLETAMENTE CALIFICADO por el Tribunal ({', '.join(lista_evals)})!")
                        st.dataframe(df_cuadro_grp, use_container_width=True)
                        
                        excel_bytes_indiv = generar_excel_cuadro_individual(df_cuadro_grp, grp_consultar, paralelo_sel)
                        st.download_button(
                            label=f"📥 Descargar Cuadro Consolidado en Excel (.xlsx) - {grp_consultar}",
                            data=excel_bytes_indiv,
                            file_name=f"Consolidado_{paralelo_sel.replace(' ', '_')}_{grp_consultar.replace(' ', '_')}.xlsx",
                            mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                            use_container_width=True,
                            key="btn_dl_transp"
                        )
                    else:
                        st.info(f"⏳ Este grupo cuenta con {cant_evals} de {total_tribunal_necesario} calificaciones: {', '.join(lista_evals)}. El cuadro se activará cuando los 3 miembros completen la evaluación.")
                else:
                    st.warning("Este grupo aún no tiene calificaciones registradas.")

    # Formulario activo de calificación
    if grupo_activo:
        with st.form("form_calificacion"):
            respuestas_escrito = []
            campos_faltantes = []
            
            # --- EVALUACIÓN SUPLETORIO VS ORDINARIO ---
            if es_supletorio:
                grupo_origen = grupo_activo.replace("Supletorio ", "")
                nota_fase1_fijada, detalles_fase1_fijada = obtener_nota_escrito_grupo_original(paralelo_sel, grupo_origen)
                
                st.markdown(f"### ⚠️ EVALUACIÓN SUPLETORIA: {grupo_activo}")
                st.info(f"""
                    📌 **RÉGIMEN SUPLETORIO:**
                    * El cursante rinde únicamente la **FASE II: Sustentación Oral, Defensa e Interrogatorio (12 Puntos)**.
                    * La calificación de la **FASE I (Trabajo Escrito)** se mantiene fija con la nota grupal original: **{nota_fase1_fijada:.2f} / 8.00 puntos**.
                """)
                subtotal_escrito_aplicar = nota_fase1_fijada
                detalles_escrito_json_aplicar = detalles_fase1_fijada
            else:
                st.markdown(f"### 📄 FASE I: DESARROLLO DEL CASO — PRODUCTO ENTREGABLE FÍSICO (8 PUNTOS)")
                st.caption(f"Calificación para: **{grupo_activo}** — Tema: *{tema_activo}*")
                
                for c_idx, crit_obj in enumerate(CRITERIOS_DETALLADOS_ESCRITO):
                    c_nombre = crit_obj["criterio"]
                    c_tot = crit_obj["total_pts"]
                    
                    with st.container():
                        st.markdown(f"#### 📌 {c_nombre} *(Puntaje Máximo: {c_tot:.2f} pts)*")
                        
                        for ind_idx, ind in enumerate(crit_obj["indicadores"]):
                            c_desc, c_sel = st.columns([3, 2])
                            with c_desc:
                                st.markdown(f"**Indicador {ind_idx + 1}:** {ind['desc']}")
                                st.caption(f"Valor Máximo: {ind['max']:.2f} pts")
                            with c_sel:
                                opc_sel = st.selectbox(
                                    f"Nivel para Ind {ind_idx+1} de {c_nombre}",
                                    OPCIONES_NIVEL,
                                    index=None,
                                    placeholder="-- Seleccione calificación --",
                                    key=f"esc_{c_idx}_{ind_idx}",
                                    label_visibility="collapsed"
                                )
                                if opc_sel is not None:
                                    pts_obtenidos = ind["valores"][opc_sel]
                                    respuestas_escrito.append(pts_obtenidos)
                                else:
                                    campos_faltantes.append(f"{c_nombre} ➔ Indicador {ind_idx + 1}")
                            st.divider()

            # --- FASE II: SUSTENTACIÓN ORAL ---
            num_alumnos = len(alumnos_activos)
            titulo_fase2 = f"### 🗣️ FASE II: SUSTENTACIÓN ORAL, DEFENSA E INTERROGATORIO (12 PUNTOS) — {num_alumnos} Cursante(s)"
            st.markdown(titulo_fase2)
            st.caption("Obligatorio calificar individualmente los 4 criterios a cada cursante.")
            
            evaluaciones_orales = []
            
            for i, alumno in enumerate(alumnos_activos):
                with st.expander(f"🎓 EVALUACIÓN INDIVIDUAL ➔ Integrante {i+1}: {alumno.upper()}", expanded=True):
                    st.markdown(f"""
                        <div style="
                            background: linear-gradient(90deg, #00204A 0%, #083370 100%);
                            padding: 14px 18px;
                            border-radius: 8px;
                            border-left: 6px solid #D97706;
                            margin-bottom: 16px;
                            box-shadow: 0 2px 6px rgba(0,0,0,0.12);
                        ">
                            <span style="
                                background-color: rgba(217, 119, 6, 0.25);
                                color: #FBBF24;
                                font-size: 0.78rem;
                                font-weight: 700;
                                letter-spacing: 1.2px;
                                text-transform: uppercase;
                                padding: 3px 8px;
                                border-radius: 4px;
                            ">
                                {'Calificación de Examen Supletorio' if es_supletorio else 'Calificación Individual del Cursante'}
                            </span>
                            <h2 style="
                                margin: 6px 0 0 0;
                                color: #FFFFFF;
                                font-size: 1.35rem;
                                font-weight: 700;
                                letter-spacing: -0.3px;
                            ">
                                👤 {alumno}
                            </h2>
                        </div>
                    """, unsafe_allow_html=True)
                    
                    al_notas = {}
                    al_completo = True
                    
                    cols_oral = st.columns(2)
                    for o_idx, (o_crit, o_peso, o_desc) in enumerate(CRITERIOS_ORAL):
                        with cols_oral[o_idx % 2]:
                            st.markdown(f"**{o_crit}** *(Máx: {o_peso:.2f} pts)*")
                            st.caption(o_desc)
                            calif_oral = st.radio(
                                f"Calif_{alumno}_{o_idx}",
                                OPCIONES_NIVEL,
                                index=None,
                                horizontal=True,
                                key=f"oral_{i}_{o_idx}",
                                label_visibility="collapsed"
                            )
                            if calif_oral is None:
                                campos_faltantes.append(f"{alumno} ➔ {o_crit}")
                                al_completo = False
                            else:
                                al_notas[o_idx] = {
                                    "nivel": calif_oral,
                                    "pts": PUNTOS_ORAL[calif_oral]
                                }
                            st.markdown("<div style='margin-bottom: 12px;'></div>", unsafe_allow_html=True)
                    
                    if al_completo:
                        subtotal_indiv = sum(n["pts"] for n in al_notas.values())
                        evaluaciones_orales.append({
                            "alumno": alumno,
                            "sust_nivel": al_notas[0]["nivel"],
                            "sust_pts": al_notas[0]["pts"],
                            "interrog_nivel": al_notas[1]["nivel"],
                            "interrog_pts": al_notas[1]["pts"],
                            "estrateg_nivel": al_notas[2]["nivel"],
                            "estrateg_pts": al_notas[2]["pts"],
                            "diferenc_nivel": al_notas[3]["nivel"],
                            "diferenc_pts": al_notas[3]["pts"],
                            "subtotal_oral": subtotal_indiv
                        })
            
            enviar = st.form_submit_button("💾 Guardar y Enviar Evaluación", use_container_width=True)
            
            if enviar:
                if not paralelo_sel:
                    st.error("⚠️ Debe seleccionar el Paralelo en la parte superior.")
                elif not evaluador_activo:
                    st.error("⚠️ Debe seleccionar su nombre como Miembro del Tribunal.")
                elif not grupo_activo:
                    st.error("⚠️ Debe seleccionar el Grupo a Evaluar.")
                elif len(campos_faltantes) > 0:
                    st.error("🚨 **EVALUACIÓN INCOMPLETA: Hay campos o indicadores sin calificar.**")
                    st.warning("Por favor complete los siguientes ítems antes de guardar:")
                    for item in campos_faltantes:
                        st.write(f"❌ **{item}**")
                else:
                    cursor = conn.cursor()
                    cursor.execute("SELECT id FROM evaluaciones WHERE paralelo = ? AND evaluador = ? AND grupo = ?", (paralelo_sel, evaluador_activo, grupo_activo))
                    if cursor.fetchone():
                        st.error("⚠️ Esta evaluación ya cuenta con su calificación registrada.")
                    else:
                        if not es_supletorio:
                            subtotal_escrito_aplicar = round(sum(respuestas_escrito), 2)
                            detalles_escrito_json_aplicar = json.dumps(respuestas_escrito)
                            
                        fecha_reg = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
                        
                        for ev in evaluaciones_orales:
                            nota_total_calc = round(subtotal_escrito_aplicar + ev["subtotal_oral"], 2)
                            cursor.execute("""
                                INSERT INTO evaluaciones 
                                (fecha, paralelo, evaluador, grupo, alumno, nota_escrito, detalles_escrito_json,
                                 oral_sustentacion_nivel, oral_sustentacion_pts,
                                 oral_interrogatorio_nivel, oral_interrogatorio_pts,
                                 oral_estrategico_nivel, oral_estrategico_pts,
                                 oral_diferenciador_nivel, oral_diferenciador_pts,
                                 nota_oral, nota_total)
                                VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)
                            """, (
                                fecha_reg, paralelo_sel, evaluador_activo, grupo_activo, ev["alumno"],
                                subtotal_escrito_aplicar, detalles_escrito_json_aplicar,
                                ev["sust_nivel"], ev["sust_pts"],
                                ev["interrog_nivel"], ev["interrog_pts"],
                                ev["estrateg_nivel"], ev["estrateg_pts"],
                                ev["diferenc_nivel"], ev["diferenc_pts"],
                                ev["subtotal_oral"], nota_total_calc
                            ))
                        conn.commit()
                        respaldar_evaluaciones_en_disco()
                        st.session_state["grupo_evaluado_exito"] = (paralelo_sel, grupo_activo)
                        st.success(f"✅ ¡Evaluación del {grupo_activo} guardada exitosamente!")
                        st.balloons()
                        st.rerun()

    if "grupo_evaluado_exito" in st.session_state:
        p_eval, g_eval = st.session_state["grupo_evaluado_exito"]
        st.info(f"📌 Última evaluación registrada por usted: **{g_eval}** ({p_eval}).")
        
        docentes_p_check, grupos_p_check, _ = cargar_configuracion(p_eval)
        df_cuadro_check, cant_evals_check, lista_evals_check = calcular_cuadro_consolidado_grupo(p_eval, g_eval, grupos_p_check)
        total_tribunal_necesario = 3
        
        if df_cuadro_check is not None and not df_cuadro_check.empty:
            if cant_evals_check >= total_tribunal_necesario:
                st.success(f"🎉 ¡TODOS LOS DOCENTES ({cant_evals_check}/{total_tribunal_necesario}) HAN CULMINADO LA CALIFICACIÓN DE {g_eval}!")
                st.markdown("#### Cuadro de Calificación Oficial:")
                st.dataframe(df_cuadro_check, use_container_width=True)
                
                excel_bytes_indiv = generar_excel_cuadro_individual(df_cuadro_check, g_eval, p_eval)
                st.download_button(
                    label=f"📥 Descargar Cuadro Oficial en Excel (.xlsx) - {g_eval}",
                    data=excel_bytes_indiv,
                    file_name=f"Consolidado_{p_eval.replace(' ', '_')}_{g_eval.replace(' ', '_')}.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                    use_container_width=True,
                    key="btn_dl_post_eval"
                )
            else:
                st.warning(f"⏳ Evaluación registrada. Han calificado {cant_evals_check} de {total_tribunal_necesario} miembros del tribunal.")

# =============================================================
# --- PESTAÑA ADMINISTRADOR ---
# =============================================================
with pestana_admin:
    st.subheader("Panel de Control del Administrador")
    password = st.text_input("Ingrese la Contraseña:", type="password")
    
    if password == "AdminPolicial2026":
        tabs_admin = st.tabs(["📊 Consolidado de Notas", "📄 Generar Actas Word (.docx)", "⚙️ Gestión de Docentes y Grupos"])
        
        with tabs_admin[0]:
            df = pd.read_sql_query("SELECT * FROM evaluaciones", conn)
            if df.empty:
                st.info("No hay evaluaciones registradas hasta el momento.")
            else:
                if "paralelo" not in df.columns:
                    df["paralelo"] = "Paralelo A"
                    
                col_filtro_p, col_filtro_g = st.columns(2)
                with col_filtro_p:
                    filtro_admin_p = st.selectbox("1. Filtrar por Paralelo:", ["Todos", "Paralelo A", "Paralelo B"])
                
                df_p = df if filtro_admin_p == "Todos" else df[df["paralelo"] == filtro_admin_p]
                
                with col_filtro_g:
                    grupos_unicos = ["Todos"] + sorted(list(df_p["grupo"].dropna().unique()))
                    filtro_admin_g = st.selectbox("2. Filtrar por Grupo:", grupos_unicos)
                    
                df_filtrado = df_p if filtro_admin_g == "Todos" else df_p[df_p["grupo"] == filtro_admin_g]
                
                titulo_tabla = f"Promedios Finales por Alumno ({filtro_admin_p} - {filtro_admin_g})"
                st.markdown(f"#### {titulo_tabla}")
                
                if not df_filtrado.empty:
                    resumen = df_filtrado.groupby(["paralelo", "grupo", "alumno"]).agg(
                        nota_escrito_prom=("nota_escrito", "mean"),
                        nota_oral_prom=("nota_oral", "mean"),
                        nota_total_prom=("nota_total", "mean")
                    ).round(2).reset_index()
                    
                    lista_excel = []
                    for idx, row in resumen.iterrows():
                        grado, nombres_apellidos = separar_grado_nombre(row["alumno"])
                        lista_excel.append({
                            "Ord.": idx + 1,
                            "Paralelo": row["paralelo"],
                            "Grupo": row["grupo"],
                            "Grado": grado,
                            "Nombres y Apellidos": nombres_apellidos,
                            "Calificación del Trabajo Escrito": row["nota_escrito_prom"],
                            "Calificación SUSTENTACIÓN ORAL INDIVIDUAL": row["nota_oral_prom"],
                            "Total Calificación Final": row["nota_total_prom"]
                        })
                    
                    df_oficial = pd.DataFrame(lista_excel)
                    st.dataframe(df_oficial, use_container_width=True)
                    
                    output_excel = io.BytesIO()
                    with pd.ExcelWriter(output_excel, engine="openpyxl") as writer:
                        df_oficial.to_excel(writer, index=False, sheet_name="Acta de Calificaciones")
                    excel_data = output_excel.getvalue()
                    
                    st.download_button(
                        label=f"📊 Descargar Resumen en Excel (.xlsx) - {filtro_admin_g}",
                        data=excel_data,
                        file_name=f"Acta_{filtro_admin_p.replace(' ', '_')}_{filtro_admin_g.replace(' ', '_')}_{datetime.now().strftime('%Y%m%d_%H%M')}.xlsx",
                        mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                        use_container_width=True
                    )
                else:
                    st.info("No hay registros aún para los filtros seleccionados.")
            
            st.divider()
            with st.expander("⚠️ Zona de Mantenimiento: Vaciar Evaluaciones de Prueba"):
                st.warning("Esta acción eliminará únicamente las calificaciones registradas de prueba. No afectará a los docentes ni a los grupos configurados.")
                confirmar = st.checkbox("Confirmo que deseo borrar todas las calificaciones de prueba registradas.")
                if st.button("🗑️ Borrar Todas las Calificaciones", disabled=not confirmar):
                    cursor = conn.cursor()
                    cursor.execute("DELETE FROM evaluaciones")
                    conn.commit()
                    if os.path.exists(ARCHIVO_RESPALDO):
                        try:
                            os.remove(ARCHIVO_RESPALDO)
                        except Exception:
                            pass
                    st.success("✅ Base de datos de calificaciones reiniciada a cero.")
                    st.rerun()

        with tabs_admin[1]:
            st.markdown("### 🖨️ Generación de Actas Físicas Oficiales (Formato Institucional Word)")
            st.caption("Consolida automáticamente las notas individuales y grupales para la emisión del acta oficial.")
            
            plantilla_oficial = "INDICADORES PARA EVALUACIÓN DEL EJERCICIO PRACTICO.docx"
            if not os.path.exists(plantilla_oficial):
                st.error(f"⚠️ No se encontró el archivo de plantilla `{plantilla_oficial}` en el repositorio.")
                st.info("Asegúrese de subir la plantilla Word con ese nombre exacto a su repositorio en GitHub.")
            else:
                par_acta = st.selectbox("Seleccione el Paralelo:", PARALELOS_DISPONIBLES, key="par_acta_sel")
                docentes_p, grupos_p, temas_p = cargar_configuracion(par_acta)
                supletorios_act = obtener_supletorios_pendientes(par_acta, grupos_p)
                grupos_para_actas = list(grupos_p.keys()) + list(supletorios_act.keys())
                
                if grupos_para_actas:
                    grp_acta = st.selectbox("Seleccione el Grupo:", grupos_para_actas, key="grp_acta_sel")
                    
                    cursor = conn.cursor()
                    cursor.execute("""
                        SELECT evaluador, alumno, detalles_escrito_json, 
                               oral_sustentacion_pts, oral_interrogatorio_pts, 
                               oral_estrategico_pts, oral_diferenciador_pts, 
                               nota_oral, fecha
                        FROM evaluaciones
                        WHERE paralelo = ? AND grupo = ?
                    """, (par_acta, grp_acta))
                    registros = cursor.fetchall()
                    
                    if registros:
                        evaluadores_registrados = list(dict.fromkeys(r[0] for r in registros))
                        st.success(f"📋 Evaluaciones registradas de {len(evaluadores_registrados)} evaluador(es): {', '.join(evaluadores_registrados)}")
                        
                        tribunal_para_acta = docentes_p[:3] if len(docentes_p) >= 3 else evaluadores_registrados
                        
                        listas_indicadores = []
                        for r in registros:
                            if r[2]:
                                try:
                                    listas_indicadores.append(json.loads(r[2]))
                                except Exception:
                                    pass
                                    
                        notas_21_prom = [0.0] * 21
                        if listas_indicadores:
                            num_listas = len(listas_indicadores)
                            for i in range(21):
                                notas_21_prom[i] = round(sum(l[i] for l in listas_indicadores if i < len(l)) / num_listas, 2)
                        
                        integrantes_acta = grupos_p.get(grp_acta, [])
                        if not integrantes_acta and grp_acta in supletorios_act:
                            integrantes_acta = supletorios_act[grp_acta]["alumnos"]
                            
                        desglose_alumnos = []
                        for al_nombre in integrantes_acta:
                            evals_al = [r for r in registros if r[1] == al_nombre]
                            if evals_al:
                                p_sust = round(sum(r[3] for r in evals_al) / len(evals_al), 2)
                                p_interrog = round(sum(r[4] for r in evals_al) / len(evals_al), 2)
                                p_estrateg = round(sum(r[5] for r in evals_al) / len(evals_al), 2)
                                p_diferenc = round(sum(r[6] for r in evals_al) / len(evals_al), 2)
                                p_tot_oral = round(sum(r[7] for r in evals_al) / len(evals_al), 2)
                            else:
                                p_sust = p_interrog = p_estrateg = p_diferenc = p_tot_oral = 0.0
                                
                            desglose_alumnos.append({
                                "nombre": al_nombre,
                                "sust": p_sust,
                                "interrog": p_interrog,
                                "estrateg": p_estrateg,
                                "diferenc": p_diferenc,
                                "total_oral": p_tot_oral
                            })
                        
                        fecha_acta = registros[0][8] if registros[0][8] else datetime.now().strftime("%d de %B de %Y")
                        tema_acta = temas_p.get(grp_acta, "Proyecto Integrador de Análisis Criminal")
                        if grp_acta in supletorios_act:
                            tema_acta = supletorios_act[grp_acta]["tema"]
                        
                        acta_bytes = generar_acta_docx_oficial(
                            plantilla_path=plantilla_oficial,
                            paralelo=par_acta,
                            grupo_nombre=grp_acta,
                            tema_caso=tema_acta,
                            fecha_str=fecha_acta,
                            integrantes=integrantes_acta,
                            tribunal_nombres=tribunal_para_acta,
                            notas_escrito_21=notas_21_prom,
                            desglose_alumnos=desglose_alumnos
                        )
                        
                        st.download_button(
                            label=f"📥 Descargar Acta Oficial en Word (.docx) - {grp_acta}",
                            data=acta_bytes,
                            file_name=f"Acta_{par_acta.replace(' ', '_')}_{grp_acta.replace(' ', '_')}.docx",
                            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                            use_container_width=True
                        )
                    else:
                        st.warning(f"Aún no hay calificaciones registradas para el {grp_acta}.")
                else:
                    st.info("No hay grupos registrados en este paralelo.")

        with tabs_admin[2]:
            st.markdown("### 👥 Configuración Adicional de Docentes y Grupos")
            paralelo_gestionar = st.selectbox("Seleccione el Paralelo a configurar:", PARALELOS_DISPONIBLES)
            
            c_docentes, c_grupos = st.columns(2)
            
            with c_docentes:
                st.markdown(f"#### 1. Miembros del Tribunal ({paralelo_gestionar})")
                with st.form("form_agregar_docente", clear_on_submit=True):
                    nuevo_docente = st.text_input("Grado y Nombre del Evaluador (ej. Mayor Patricio Ramos):")
                    btn_docente = st.form_submit_button("➕ Agregar Evaluador")
                    if btn_docente and nuevo_docente.strip():
                        cursor = conn.cursor()
                        cursor.execute("INSERT INTO config_docentes (paralelo, nombre) VALUES (?, ?)", (paralelo_gestionar, nuevo_docente.strip()))
                        conn.commit()
                        st.success(f"Evaluador agregado: {nuevo_docente.strip()}")
                        st.rerun()
                
                cursor = conn.cursor()
                cursor.execute("SELECT id, nombre FROM config_docentes WHERE paralelo = ?", (paralelo_gestionar,))
                docs_actuales = cursor.fetchall()
                if docs_actuales:
                    st.markdown("**Evaluadores registrados:**")
                    for d_id, d_nom in docs_actuales:
                        col_d1, col_d2 = st.columns([4, 1])
                        col_d1.write(f"• {d_nom}")
                        if col_d2.button("🗑️", key=f"del_doc_{d_id}"):
                            cursor.execute("DELETE FROM config_docentes WHERE id = ?", (d_id,))
                            conn.commit()
                            st.rerun()
                else:
                    st.caption("No hay evaluadores agregados para este paralelo.")
                    
            with c_grupos:
                st.markdown(f"#### 2. Grupos, Tema de Caso y Alumnos ({paralelo_gestionar})")
                with st.form("form_agregar_grupo", clear_on_submit=True):
                    nom_grupo = st.text_input("Identificador del Grupo (ej. Grupo 1):")
                    tema_grupo = st.text_input("Tema del Caso Analizado (ej. Análisis de Redes de Vínculos):")
                    alumnos_texto = st.text_area(
                        "Integrantes (un nombre con grado por línea):",
                        placeholder="Cptn. Garzon Masapanta Henry Rodrigo\nSgos. Cupueran Vallejos Juan Diego"
                    )
                    btn_grupo = st.form_submit_button("➕ Registrar Grupo y Alumnos")
                    
                    if btn_grupo:
                        alumnos_lista = [linea.strip() for linea in alumnos_texto.split("\n") if linea.strip()]
                        if not nom_grupo.strip():
                            st.error("Debe ingresar un nombre para el grupo.")
                        elif len(alumnos_lista) == 0:
                            st.error("Debe ingresar al menos un alumno para este grupo.")
                        else:
                            cursor = conn.cursor()
                            cursor.execute(
                                "INSERT INTO config_grupos (paralelo, grupo, tema_caso, alumnos_json) VALUES (?, ?, ?, ?)",
                                (paralelo_gestionar, nom_grupo.strip(), tema_grupo.strip(), json.dumps(alumnos_lista))
                            )
                            conn.commit()
                            st.success(f"Grupo registrado exitosamente con {len(alumnos_lista)} integrantes.")
                            st.rerun()
                
                cursor.execute("SELECT id, grupo, alumnos_json, tema_caso FROM config_grupos WHERE paralelo = ?", (paralelo_gestionar,))
                grps_actuales = cursor.fetchall()
                if grps_actuales:
                    st.markdown("**Grupos registrados:**")
                    for g_id, g_nom, g_alums_raw, g_tema in grps_actuales:
                        alums = json.loads(g_alums_raw)
                        with st.expander(f"📁 {g_nom} ({len(alums)} integrantes) - Tema: {g_tema}"):
                            texto_actual = "\n".join(alums)
                            nuevo_tema = st.text_input("Tema del Caso:", value=g_tema, key=f"edit_tema_{g_id}")
                            nuevo_texto_alums = st.text_area(
                                "Modificar integrantes (un nombre por línea):",
                                value=texto_actual,
                                key=f"edit_txt_{g_id}"
                            )
                            col_guardar, col_borrar = st.columns(2)
                            if col_guardar.button("💾 Guardar Cambios", key=f"btn_save_{g_id}"):
                                nueva_lista = [l.strip() for l in nuevo_texto_alums.split("\n") if l.strip()]
                                cursor.execute(
                                    "UPDATE config_grupos SET alumnos_json = ?, tema_caso = ? WHERE id = ?",
                                    (json.dumps(nueva_lista), nuevo_tema.strip(), g_id)
                                )
                                conn.commit()
                                st.success("Datos actualizados.")
                                st.rerun()
                            if col_borrar.button("🗑️ Eliminar Grupo", key=f"del_grp_{g_id}"):
                                cursor.execute("DELETE FROM config_grupos WHERE id = ?", (g_id,))
                                conn.commit()
                                st.rerun()
                else:
                    st.caption("No hay grupos agregados para este paralelo.")
                    
    elif password != "":
        st.error("Contraseña incorrecta.")
